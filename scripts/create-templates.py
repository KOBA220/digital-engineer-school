from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]/'templates'
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
navy=RGBColor.from_string('172B3A');green=RGBColor.from_string('214D43');grey=RGBColor.from_string('687C89')
def tb(s,x,y,w,h,text,size=22,color=navy,bold=False):
 box=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=box.text_frame;tf.word_wrap=True
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Noto Sans JP';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color;p.space_after=Pt(12)
 return box
slides=[('プロジェクト報告','[プロジェクト名]\n[担当部署・発表者]\n[日付]'),('今週の進捗','[今回伝えたい結論を1文で]\n\n完了したこと\n[成果物・確認結果]\n\n次に進めること\n[担当者・期限]'),('相談したい課題','[判断してほしいこと]\n\n現状と影響\n[課題が起きる条件・影響する範囲]\n\n提案と依頼\n[対応案・必要な支援・判断期限]'),('システム構成','[構成図・画面イメージをここに配置]\n\n利用者　[対象ユーザー]\n処理　　[入力から結果までの流れ]\n連携　　[外部サービス・保存先]'),('スプリントの結果','[今回のスプリント名]\n\n目標　[スプリントの目標]\n結果　[達成したこと・未完了のこと]\n改善　[次回変えること]')]
for idx,(title,body) in enumerate(slides):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('F8FAFB')
 tb(s,.7,.55,11.9,.75,title,32,green,True);tb(s,.75,1.65,11.8,4.9,body,23);tb(s,.75,6.85,10,.3,'デジタルエンジニア学校',11,grey);tb(s,12,6.85,.6,.3,str(idx+1),11,grey)
s=prs.slides.add_slide(prs.slide_layouts[6]);tb(s,.7,.55,11.9,.75,'開発ロードマップ',32,green,True)
table=s.shapes.add_table(5,5,Inches(.75),Inches(1.7),Inches(11.8),Inches(3.8)).table
rows=[['工程','第1週','第2週','第3週','第4週'],['要件・設計','[予定]','','',''],['実装','','[予定]','[予定]',''],['検証','','','[予定]',''],['公開・振り返り','','','','[予定]']]
for r,row in enumerate(rows):
 for c,text in enumerate(row):
  cell=table.cell(r,c);cell.text=text;cell.fill.solid();cell.fill.fore_color.rgb=green if r==0 else RGBColor.from_string('EAF0ED' if c==0 else 'FFFFFF')
  for p in cell.text_frame.paragraphs:p.font.name='Noto Sans JP';p.font.size=Pt(17);p.font.color.rgb=RGBColor.from_string('FFFFFF') if r==0 else navy
# Editable table rather than an illustrative bitmap.
tb(s,.75,5.9,11.5,.6,'[各工程の担当者と完了条件を記入]',20,grey);tb(s,.75,6.85,10,.3,'デジタルエンジニア学校',11,grey)
prs.save(root/'school-presentation.pptx')
doc=Document();sec=doc.sections[0];sec.top_margin=DInches(.75);sec.bottom_margin=DInches(.75);sec.left_margin=DInches(.8);sec.right_margin=DInches(.8)
styles=doc.styles
for border in list(styles.element.findall('.//' + qn('w:pBdr'))):
 border.getparent().remove(border)
for name in ['Normal','Title','Heading 1','Heading 2']:
 styles[name].font.name='Noto Sans JP';styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans JP')
styles['Normal'].font.size=DPt(10.5);styles['Normal'].paragraph_format.space_after=DPt(8)
for name in ['Title','Heading 1','Heading 2']:styles[name].font.color.rgb=DColor.from_string('214D43')
doc.add_heading('システム仕様書',0);doc.add_paragraph('[プロジェクト名]');doc.add_paragraph('版：[1.0]　更新日：[YYYY-MM-DD]　担当部署：[部署名]')
sections=[('1. 概要・対象範囲',['項目','記入内容'],[['対象ユーザー','[誰が利用するか]'],['対象業務','[何を支援するか]'],['対象範囲','[今回実装する範囲]'],['対象外','[今回実装しない範囲]']]),('2. 機能・画面',['ID','機能・画面','入力・操作','結果'],[['F-001','[機能名]','[入力項目・操作]','[保存・表示・通知]'],['F-002','[機能名]','[入力項目・操作]','[保存・表示・通知]']]),('3. 権限・公開範囲',['対象','閲覧できる人','編集できる人'],[['Privateプロジェクト','[招待メンバー]','[編集者・所有者]'],['Publicプロジェクト','[学校内ユーザー]','[編集者・所有者]']]),('4. API・データ',['項目','仕様'],[['API','[URL・メソッド・認証・リクエスト・レスポンス]'],['保存データ','[項目・型・必須・保持期限]'],['外部連携','[接続先・権限・費用・障害時の扱い]']]),('5. 受入条件',['ID','操作・条件','期待結果'],[['AC-001','[操作・前提条件]','[成功時の結果]'],['AC-002','[権限なし・入力不備]','[拒否・エラー表示]']]),('6. 運用・変更履歴',['項目','内容'],[['運用','[バックアップ・監視・問い合わせ先]'],['変更履歴','[日付・変更点・担当者]']])]
for index,(title,headers,rows) in enumerate(sections):
 if index in [2,4]:doc.add_page_break()
 doc.add_heading(title,1);t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=v
 doc.add_paragraph('')
doc.save(root/'system-specification.docx')
print('Created 6-slide presentation and 3-page specification template')
